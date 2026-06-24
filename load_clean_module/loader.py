import os
import re
from pathlib import Path

import cv2
import numpy as np
import openpyxl

from database import DBManager
from mpivr20_cms import get_clean_output_dir

RAW_DIR = "data/raw"
PROCESSED_DIR = get_clean_output_dir()
SHEET_IMAGE = "主要"
SHEET_LABEL = "體積面積量測"
TARGET_COL = 2
DEFAULT_MAX_GROUND_TRUTH_RATIO = 0.5
DEFAULT_MIN_GROUND_TRUTH_PIXELS = 30
STORAGE_LAYOUT_TYPED_FLAT = "typed_flat_v1"
OUTPUT_TYPES = ("image", "label", "ground_truth", "mask_view")
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


class GroundTruthAreaError(ValueError):
    """Raised when all generated ground-truth masks fail quality rules."""

    def __init__(self, excel_path: str, violations: list[dict]):
        self.excel_path = excel_path
        self.violations = violations
        details = "; ".join(
            f"{item['sample_name']} {item['reason']} "
            f"({item['pixels']} px, {item['ratio']:.1%})"
            for item in violations
        )
        super().__init__(
            f"ground_truth 全部不符合清洗規則，不可用於訓練：{excel_path}；{details}"
        )


def imwrite_unicode(path: str, img: np.ndarray):
    ext = os.path.splitext(path)[1] or ".png"
    ok, buf = cv2.imencode(ext, img)
    if not ok:
        raise IOError(f"cv2.imencode 失敗：{path}")
    buf.tofile(path)


def decode_image(data: bytes) -> np.ndarray:
    arr = np.asarray(bytearray(data), dtype=np.uint8)
    return cv2.imdecode(arr, cv2.IMREAD_COLOR)


def find_excel(product_dir: str):
    for filename in os.listdir(product_dir):
        if filename.startswith("~$"):
            continue
        if filename.lower().endswith((".xlsx", ".xlsm")):
            return os.path.join(product_dir, filename)
    return None


def extract_from_sheet(sheet, key: str, data: dict):
    for img in sheet._images:
        row = img.anchor._from.row + 1
        col = img.anchor._from.col + 1
        if col != TARGET_COL:
            continue
        name = sheet.cell(row, 1).value
        if name is None:
            continue
        data.setdefault(str(name), {})[key] = decode_image(img._data())


def ground_truth_ratio(ground_truth: np.ndarray) -> float:
    """Return the non-zero mask ratio of a generated ground-truth image."""
    if ground_truth.size == 0:
        return 0.0
    return float(np.count_nonzero(ground_truth)) / float(ground_truth.size)


def ground_truth_violation(
    sample_name: str,
    ground_truth: np.ndarray,
    max_ratio: float = DEFAULT_MAX_GROUND_TRUTH_RATIO,
    min_pixels: int = DEFAULT_MIN_GROUND_TRUTH_PIXELS,
) -> dict | None:
    """Return violation details when ground-truth area is outside quality limits."""
    pixels = int(np.count_nonzero(ground_truth))
    ratio = ground_truth_ratio(ground_truth)
    if ratio > max_ratio:
        return {
            "sample_name": sample_name,
            "pixels": pixels,
            "ratio": ratio,
            "reason": "area_too_large",
            "limit": max_ratio,
        }
    if pixels < min_pixels:
        return {
            "sample_name": sample_name,
            "pixels": pixels,
            "ratio": ratio,
            "reason": "area_too_small",
            "limit": min_pixels,
        }
    return None


def violation_message(violation: dict) -> str:
    """Format a ground-truth quality violation for logs and UI."""
    if violation["reason"] == "area_too_large":
        return f"ground_truth {violation['ratio']:.1%} > {violation['limit']:.0%}"
    if violation["reason"] == "area_too_small":
        return f"ground_truth {violation['pixels']} px < {violation['limit']} px"
    return "ground_truth 不符合清洗規則"


def safe_filename_part(value: object) -> str:
    """Return a Windows-safe filename part without changing readable text."""
    text = str(value).strip()
    text = re.sub(r'[<>:"/\\\\|?*\x00-\x1f]+', "_", text)
    text = re.sub(r"\s+", "_", text)
    text = text.strip(" ._")
    if not text:
        text = "unnamed"
    if text.upper() in WINDOWS_RESERVED_NAMES:
        text = f"_{text}"
    return text


class LoadCleanService:
    def __init__(
        self,
        tasks=None,
        save_dir=PROCESSED_DIR,
        force=False,
        db=None,
        max_ground_truth_ratio=DEFAULT_MAX_GROUND_TRUTH_RATIO,
        min_ground_truth_pixels=DEFAULT_MIN_GROUND_TRUTH_PIXELS,
    ):
        self.tasks = tasks or []
        self.save_dir = save_dir
        self.force = force
        self.db = db or DBManager()
        self.max_ground_truth_ratio = max_ground_truth_ratio
        self.min_ground_truth_pixels = min_ground_truth_pixels

    @staticmethod
    def scan_tasks(raw_dir=RAW_DIR):
        if not os.path.isdir(raw_dir):
            raise FileNotFoundError(f"找不到 {raw_dir}")
        tasks = []
        for product in sorted(os.listdir(raw_dir)):
            product_dir = os.path.join(raw_dir, product)
            if not os.path.isdir(product_dir):
                continue
            excel_path = find_excel(product_dir)
            if excel_path:
                tasks.append({"excel": excel_path, "output_name": product})
        return tasks

    def run(self):
        self.db.init_db()
        if not self.tasks:
            self.tasks = self.scan_tasks()
        if not self.tasks:
            print("沒有可處理的 Excel 任務")
            return {"dataset_ids": [], "total_samples": 0, "failed": [], "failed_samples": []}

        dataset_ids = []
        total_samples = 0
        failed = []
        failed_samples = []
        for task in self.tasks:
            excel_path = task["excel"]
            output_name = task["output_name"]
            try:
                dataset_id, n_samples, reused, sample_failures = self.process_task(
                    excel_path, output_name
                )
            except Exception as exc:
                failed.append(
                    {
                        "excel": excel_path,
                        "output_name": output_name,
                        "error": f"{type(exc).__name__}: {exc}",
                    }
                )
                print(f"[{output_name}] 清洗失敗，已略過：{excel_path}")
                print(f"  {type(exc).__name__}: {exc}")
                continue
            dataset_ids.append(dataset_id)
            total_samples += n_samples
            action = "沿用快取" if reused else "完成清洗"
            print(f"[{output_name}] {action}：{n_samples} samples, dataset_id={dataset_id}")
            for failure in sample_failures:
                failed_samples.append(
                    dict(failure, excel=excel_path, output_name=output_name)
                )
                print(
                    f"  sample 標記 failed：{failure['sample_name']} "
                    f"{violation_message(failure)}"
                )

        print(
            f"全部完成：成功 {len(dataset_ids)} dataset，"
            f"{total_samples} samples，略過 {len(failed)} 檔，"
            f"failed samples {len(failed_samples)} 筆"
        )
        return {
            "dataset_ids": dataset_ids,
            "total_samples": total_samples,
            "failed": failed,
            "failed_samples": failed_samples,
        }

    def process_task(self, excel_path: str, output_name: str):
        if not os.path.isfile(excel_path):
            raise FileNotFoundError(f"檔案不存在：{excel_path}")

        source_id = self.db.upsert_source_excel(excel_path, output_name)
        out_dir = os.path.abspath(self.save_dir)
        existing = self.db.get_dataset_for_excel_path(
            excel_path,
            storage_layout=STORAGE_LAYOUT_TYPED_FLAT,
        )
        if not self.force and self.db.cache_valid(existing):
            violations = self.validate_cached_ground_truth(existing)
            if violations:
                self.db.mark_dataset_failed(existing["id"])
                raise GroundTruthAreaError(excel_path, violations)
            return existing["id"], existing["n_samples"], True, []

        dataset_id = self.db.upsert_dataset(
            dataset_name=output_name,
            source_excel_id=source_id,
            processed_dir=out_dir,
            storage_layout=STORAGE_LAYOUT_TYPED_FLAT,
            status="running",
        )
        try:
            samples, area_violations = self._process_excel(excel_path, out_dir)
            active_count = self.db.replace_samples(dataset_id, samples)
            if active_count <= 0 and area_violations:
                raise GroundTruthAreaError(excel_path, area_violations)
        except Exception:
            self.db.mark_dataset_failed(dataset_id)
            raise
        return dataset_id, active_count, False, area_violations

    def validate_cached_ground_truth(self, dataset: dict) -> list[dict]:
        """Check cached ground-truth files before reusing a done dataset."""
        violations = []
        for sample in self.db.list_samples(dataset["id"]):
            ground_truth = cv2.imread(sample["ground_truth_path"], cv2.IMREAD_UNCHANGED)
            if ground_truth is None:
                raise FileNotFoundError(f"無法讀取 ground_truth：{sample['ground_truth_path']}")
            violation = ground_truth_violation(
                sample["sample_name"],
                ground_truth,
                max_ratio=self.max_ground_truth_ratio,
                min_pixels=self.min_ground_truth_pixels,
            )
            if violation:
                violation["ground_truth_path"] = sample["ground_truth_path"]
                violations.append(violation)
        return violations

    def _process_excel(self, excel_path: str, out_dir: str):
        print(f"載入 {excel_path}")
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        missing = [s for s in (SHEET_IMAGE, SHEET_LABEL) if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"缺少工作表 {missing}，現有：{wb.sheetnames}")

        data = {}
        extract_from_sheet(wb[SHEET_IMAGE], "image", data)
        extract_from_sheet(wb[SHEET_LABEL], "label", data)

        samples = []
        area_violations = []
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        output_dirs = {
            output_type: os.path.abspath(os.path.join(out_dir, output_type))
            for output_type in OUTPUT_TYPES
        }
        for output_type_dir in output_dirs.values():
            os.makedirs(output_type_dir, exist_ok=True)
        excel_name = safe_filename_part(Path(excel_path).stem)
        used_filenames = set()
        for sample_name, items in sorted(data.items()):
            if "image" not in items or "label" not in items:
                print(f"  缺資料跳過：{sample_name}")
                continue

            image = items["image"]
            label = items["label"]
            if image.shape != label.shape:
                height, width = image.shape[:2]
                label = cv2.resize(label, (width, height), interpolation=cv2.INTER_NEAREST)

            diff = cv2.absdiff(image, label)
            gray = cv2.cvtColor(diff, cv2.COLOR_BGR2GRAY)
            ground_truth = cv2.threshold(gray, 1, 1, cv2.THRESH_BINARY)[1]
            mask_view = cv2.threshold(gray, 1, 255, cv2.THRESH_BINARY)[1]
            violation = ground_truth_violation(
                str(sample_name),
                ground_truth,
                max_ratio=self.max_ground_truth_ratio,
                min_pixels=self.min_ground_truth_pixels,
            )

            filename = f"{excel_name}_{safe_filename_part(sample_name)}.png"
            if filename in used_filenames:
                raise ValueError(f"清洗輸出檔名重複，請檢查 Excel 欄位名稱：{filename}")
            used_filenames.add(filename)
            image_path = os.path.abspath(os.path.join(output_dirs["image"], filename))
            ground_truth_path = os.path.abspath(
                os.path.join(output_dirs["ground_truth"], filename)
            )
            mask_view_path = os.path.abspath(os.path.join(output_dirs["mask_view"], filename))
            label_path = os.path.abspath(os.path.join(output_dirs["label"], filename))

            imwrite_unicode(image_path, image)
            imwrite_unicode(label_path, label)
            imwrite_unicode(ground_truth_path, ground_truth)
            imwrite_unicode(mask_view_path, mask_view)
            if violation:
                violation["ground_truth_path"] = ground_truth_path
                area_violations.append(violation)
                print(
                    f"  ground_truth 不符合清洗規則：{sample_name} "
                    f"{violation_message(violation)}"
                )

            height, width = image.shape[:2]
            samples.append(
                {
                    "sample_name": str(sample_name),
                    "image_path": image_path,
                    "ground_truth_path": ground_truth_path,
                    "label_path": label_path,
                    "mask_view_path": mask_view_path,
                    "width": int(width),
                    "height": int(height),
                    "status": "failed" if violation else "active",
                }
            )

        if not samples:
            raise ValueError("沒有完整樣本可輸出")
        return samples, area_violations
