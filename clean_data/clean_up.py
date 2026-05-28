"""資料清洗：從 Excel 提取 image / label，輸出 mask 至指定資料夾。

兩種使用模式：
1. --tasks-json <file>：處理指定的 (excel, output_name) 任務清單（GUI 模式）
2. 無參數：掃描 data/raw/{product}/*.xlsx，輸出至 data/processed/{product}/
"""

import argparse
import json
import os
import sys

import cv2
import numpy as np
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

RAW_DIR = "data/raw"
PROCESSED_DIR = "data/processed"
SHEET_IMAGE = "主要"
SHEET_LABEL = "體積面積量測"
TARGET_COL = 2  # 雷射＋彩色欄位


def imwrite_unicode(path: str, img: np.ndarray):
    """cv2.imwrite 在 Windows 對中文路徑會靜默失敗，改用 imencode + tofile。"""
    ext = os.path.splitext(path)[1] or ".png"
    ok, buf = cv2.imencode(ext, img)
    if not ok:
        raise IOError(f"cv2.imencode 失敗：{path}")
    buf.tofile(path)


def decode_image(data: bytes) -> np.ndarray:
    arr = np.asarray(bytearray(data), dtype=np.uint8)
    return cv2.imdecode(arr, cv2.IMREAD_COLOR)


def find_excel(product_dir: str):
    for f in os.listdir(product_dir):
        if f.startswith("~$"):
            continue
        if f.lower().endswith((".xlsx", ".xlsm")):
            return os.path.join(product_dir, f)
    return None


def extract_from_sheet(sheet, key: str, data: dict):
    """只取 TARGET_COL 那一欄的圖片，避開高度/3D/調色板。"""
    for img in sheet._images:
        r = img.anchor._from.row + 1
        c = img.anchor._from.col + 1
        if c != TARGET_COL:
            continue
        name = sheet.cell(r, 1).value
        if name is None:
            continue
        data.setdefault(name, {})[key] = decode_image(img._data())


def process_excel(excel_path: str, output_name: str, save_dir: str) -> int:
    """處理單一 Excel，輸出至 save_dir/output_name/。回傳成功樣本數。"""
    if not os.path.isfile(excel_path):
        print(f"[{output_name}] 檔案不存在：{excel_path}")
        return 0

    print(f"[{output_name}] 載入 {os.path.basename(excel_path)}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    missing = [s for s in (SHEET_IMAGE, SHEET_LABEL) if s not in wb.sheetnames]
    if missing:
        print(f"  缺少工作表 {missing}，跳過（現有：{wb.sheetnames}）")
        return 0

    data = {}
    extract_from_sheet(wb[SHEET_IMAGE], "image", data)
    extract_from_sheet(wb[SHEET_LABEL], "label", data)

    incomplete = [k for k, v in data.items() if "image" not in v or "label" not in v]
    for k in incomplete:
        print(f"  缺資料跳過：{k}")
        del data[k]

    if not data:
        print("  沒有完整樣本可輸出")
        return 0

    for name, items in data.items():
        if items["image"].shape != items["label"].shape:
            h, w = items["image"].shape[:2]
            items["label"] = cv2.resize(
                items["label"], (w, h), interpolation=cv2.INTER_NEAREST
            )
        diff = cv2.absdiff(items["image"], items["label"])
        gray = cv2.cvtColor(diff, cv2.COLOR_BGR2GRAY)
        items["mask"] = cv2.threshold(gray, 1, 1, cv2.THRESH_BINARY)[1]
        items["mask_view"] = cv2.threshold(gray, 1, 255, cv2.THRESH_BINARY)[1]

    out_dir = os.path.join(save_dir, output_name)
    os.makedirs(out_dir, exist_ok=True)
    for name, items in data.items():
        d = os.path.join(out_dir, str(name))
        os.makedirs(d, exist_ok=True)
        for tag, img in items.items():
            imwrite_unicode(os.path.join(d, f"{tag}.png"), img)

    print(f"  完成 {len(data)} 個樣本 -> {out_dir}")
    return len(data)


def process_product(product_dir: str, save_dir: str) -> int:
    product_id = os.path.basename(product_dir)
    excel_path = find_excel(product_dir)
    if excel_path is None:
        print(f"[{product_id}] 找不到 Excel，跳過")
        return 0
    return process_excel(excel_path, product_id, save_dir)


def run_tasks(tasks_json: str, save_dir: str):
    with open(tasks_json, encoding="utf-8") as f:
        tasks = json.load(f)
    if not tasks:
        print("任務清單為空")
        return
    print(f"開始處理 {len(tasks)} 個任務...\n")
    total = 0
    for t in tasks:
        total += process_excel(t["excel"], t["output_name"], save_dir)
        print()
    print(f"全部完成，共輸出 {total} 個樣本至 {save_dir}")


def run_scan(save_dir: str):
    if not os.path.isdir(RAW_DIR):
        print(f"找不到 {RAW_DIR}（請從 repo 根目錄執行）")
        sys.exit(1)
    products = sorted(
        d for d in os.listdir(RAW_DIR) if os.path.isdir(os.path.join(RAW_DIR, d))
    )
    if not products:
        print(f"{RAW_DIR} 下沒有任何產品資料夾")
        return
    print(f"開始處理 {len(products)} 個產品...\n")
    total = 0
    for p in products:
        total += process_product(os.path.join(RAW_DIR, p), save_dir)
        print()
    print(f"全部完成，共輸出 {total} 個樣本至 {save_dir}")


def main():
    parser = argparse.ArgumentParser(description="Probe-Mark 資料清洗")
    parser.add_argument(
        "--tasks-json",
        help='JSON 任務清單，格式：[{"excel": "...", "output_name": "..."}, ...]',
    )
    parser.add_argument(
        "--save-dir",
        default=PROCESSED_DIR,
        help=f"輸出根目錄（預設：{PROCESSED_DIR}）",
    )
    args = parser.parse_args()

    if args.tasks_json:
        run_tasks(args.tasks_json, args.save_dir)
    else:
        run_scan(args.save_dir)


if __name__ == "__main__":
    main()
