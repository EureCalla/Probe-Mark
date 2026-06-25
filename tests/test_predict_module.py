import cv2
import numpy as np
from openpyxl import load_workbook

import predict_module.predictor as predictor_module
from predict_module.predictor import Predictor


def write_png(path, image):
    ok, encoded = cv2.imencode(".png", image)
    assert ok
    encoded.tofile(path)


def test_compute_iou_from_prediction_and_ground_truth_masks(tmp_path):
    pred = np.array([[255, 0], [255, 0]], dtype=np.uint8)
    target = np.array([[255, 255], [0, 0]], dtype=np.uint8)
    unicode_dir = tmp_path / "中文路徑"
    unicode_dir.mkdir()
    pred_path = unicode_dir / "預測_mask.png"
    target_path = unicode_dir / "答案_ground_truth.png"
    write_png(pred_path, pred)
    write_png(target_path, target)

    assert abs(Predictor._compute_iou(str(pred_path), str(target_path)) - (1 / 3)) < 1e-9


def test_write_db_import_summary_excel(tmp_path):
    samples = [
        {
            "id": 101,
            "sample_name": "A01",
            "dataset_id": 7,
            "dataset_name": "SourceA",
            "image_path": str(tmp_path / "image" / "SourceA_A01.png"),
            "ground_truth_path": str(tmp_path / "ground_truth" / "SourceA_A01.png"),
        },
        {
            "id": 102,
            "sample_name": "A02",
            "dataset_id": 7,
            "dataset_name": "SourceA",
            "image_path": str(tmp_path / "image" / "SourceA_A02.png"),
            "ground_truth_path": str(tmp_path / "ground_truth" / "SourceA_A02.png"),
        },
    ]
    outputs = [
        {
            "image_path": samples[0]["image_path"],
            "mask_path": str(tmp_path / "SourceA_A01_mask.png"),
            "overlay_path": str(tmp_path / "SourceA_A01_overlay.png"),
            "compare_path": str(tmp_path / "SourceA_A01_compare.png"),
            "iou": np.float64(0.75),
            "status": "done",
            "error_message": None,
            "measure": {
                "area_px": np.float64(12.0),
                "h_feret_px": np.float64(4.0),
                "v_feret_px": np.float64(3.0),
            },
        },
        {
            "image_path": samples[1]["image_path"],
            "mask_path": "",
            "overlay_path": "",
            "compare_path": "",
            "iou": None,
            "status": "failed",
            "error_message": "RuntimeError: boom",
            "measure": None,
        },
    ]

    summary_path = Predictor._write_db_import_summary(tmp_path, 55, samples, outputs)

    wb = load_workbook(summary_path, data_only=True)
    ws = wb["test_summary"]
    headers = [cell.value for cell in ws[1]]
    row = {header: ws.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}

    assert row["sample_id"] == 101
    assert row["image_filename"] == "SourceA_A01.png"
    assert row["status"] == "done"
    assert row["error_message"] is None
    assert row["iou"] == 0.75
    assert row["area_px"] == 12
    assert row["h_feret_px"] == 4
    assert row["v_feret_px"] == 3
    failed_row = {
        header: ws.cell(row=3, column=index + 1).value
        for index, header in enumerate(headers)
    }
    assert failed_row["sample_id"] == 102
    assert failed_row["status"] == "failed"
    assert failed_row["error_message"] == "RuntimeError: boom"
    assert failed_row["iou"] is None
    assert wb["metadata"]["B1"].value == 55
    assert wb["metadata"]["B2"].value == 2
    assert wb["metadata"]["B3"].value == 1
    assert wb["metadata"]["B4"].value == 1
    assert wb["metadata"]["B5"].value == 0.75


def test_database_predict_skips_failed_sample_and_keeps_metrics(tmp_path, monkeypatch):
    model_path = tmp_path / "best_model.pt"
    model_path.write_bytes(b"fake")
    output_dir = tmp_path / "predictions"
    gt = np.array([[255, 0], [0, 0]], dtype=np.uint8)
    gt_ok = tmp_path / "gt_ok.png"
    gt_bad = tmp_path / "gt_bad.png"
    write_png(gt_ok, gt)
    write_png(gt_bad, gt)

    samples = [
        {
            "id": 1,
            "sample_name": "ok",
            "dataset_id": 10,
            "dataset_name": "demo",
            "image_path": str(tmp_path / "ok.png"),
            "ground_truth_path": str(gt_ok),
        },
        {
            "id": 2,
            "sample_name": "bad",
            "dataset_id": 10,
            "dataset_name": "demo",
            "image_path": str(tmp_path / "bad.png"),
            "ground_truth_path": str(gt_bad),
        },
    ]

    class FakeParam:
        def numel(self):
            return 5

    class FakeModel:
        def parameters(self):
            return [FakeParam()]

    class FakeCore:
        def __init__(self, opt):
            self.output_dir = opt.output_dir
            self.model = FakeModel()

        def predict(self, image_path):
            if image_path.endswith("bad.png"):
                raise RuntimeError("boom")
            mask_path = output_dir / "ok_mask.png"
            overlay_path = output_dir / "ok_overlay.png"
            compare_path = output_dir / "ok_compare.png"
            write_png(mask_path, gt)
            write_png(overlay_path, gt)
            write_png(compare_path, gt)
            return {
                "mask_path": str(mask_path),
                "overlay_path": str(overlay_path),
                "compare_path": str(compare_path),
                "measure": {"area_px": 1.0, "h_feret_px": 1.0, "v_feret_px": 1.0},
            }

    class FakeDB:
        def __init__(self):
            self.outputs = []
            self.run_status = None
            self.metrics = None

        def init_db(self):
            pass

        def get_model(self, model_id):
            return {
                "id": model_id,
                "best_model_path": str(model_path),
                "decoder_name": "FPN",
                "encoder_name": "resnet18",
                "parameter_count": None,
            }

        def list_db_import_samples_for_model(self, model_id):
            return samples

        def insert_prediction_run(self, model_id, input_path, output_dir, mode):
            return 77

        def update_model_parameter_count(self, model_id, parameter_count):
            self.parameter_count = parameter_count

        def insert_prediction_output(self, **kwargs):
            self.outputs.append(kwargs)

        def update_prediction_run(self, prediction_run_id, status):
            self.run_status = status

        def update_model_db_import_metrics(self, model_id, mean_iou, n_samples, prediction_run_id):
            self.metrics = {
                "mean_iou": mean_iou,
                "n_samples": n_samples,
                "prediction_run_id": prediction_run_id,
            }

    fake_db = FakeDB()
    monkeypatch.setattr(predictor_module, "_core_predictor_class", lambda: FakeCore)

    result = Predictor(
        model_id=1,
        output_dir=str(output_dir),
        input_mode="database",
        db=fake_db,
    ).run()

    assert fake_db.run_status == "done"
    assert fake_db.metrics == {
        "mean_iou": 1.0,
        "n_samples": 1,
        "prediction_run_id": 77,
    }
    assert [output["status"] for output in result["outputs"]] == ["done", "failed"]
    assert fake_db.outputs[1]["status"] == "failed"
    assert "RuntimeError: boom" == fake_db.outputs[1]["error_message"]
    wb = load_workbook(result["summary_path"], data_only=True)
    ws = wb["test_summary"]
    headers = [cell.value for cell in ws[1]]
    failed_row = {
        header: ws.cell(row=3, column=index + 1).value
        for index, header in enumerate(headers)
    }
    assert failed_row["sample_id"] == 2
    assert failed_row["status"] == "failed"
    assert failed_row["error_message"] == "RuntimeError: boom"
